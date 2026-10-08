"""Build slides-reference.pptx, the template for the slides' PowerPoint output.

Starts from pandoc's default reference.pptx and gives it the look of the
HTML slides: Arial, black text, left-aligned bold titles, a title slide with
a McGill red bar beside the title and subtitle and the logos from logos/
along the bottom, and a large content area. slides-pptx.qmd uses three of
its layouts: Title Slide, Title and Content (lists and tables on their own)
and Content with Caption (a table or figure with a short note, which goes
under it in small grey text).

Run from reports-src/ after changing it:
    python build_slides_reference.py
Needs Quarto (for pandoc's default template) and python-pptx.
"""
import io
import re
import subprocess
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.shapes.picture import CT_Picture
from pptx.util import Emu

OUTPUT = "slides-reference.pptx"
INK, SOFT, MUTED = "1A1A1A", "4A4A4A", "5C5C5C"
RED, TEAL = "ED1B2F", "1A6C7A"  # McGill red, StoP-AD teal (docs/STYLE.md)
# PNG, as PowerPoint shows SVG only from 2016 on; mcgill.png is the reversed
# wordmark on its red box
LOGOS = ["logos/mcgill.png", "logos/douglas.png", "logos/stopad.png"]
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Geometry in EMU on pandoc's 16:9 slide (10 x 5.625 in)
W, H = 9144000, 5143500
MARGIN = 457200
TITLE_TOP, TITLE_H = 200000, 560000
BODY_TOP = TITLE_TOP + TITLE_H + 190000
BODY_H = H - BODY_TOP - 250000
NOTE_TOP, NOTE_H = H - 1050000, 800000
# Title slide: red bar, then title and subtitle; logos along the bottom
BAR_X, TEXT_X = MARGIN, MARGIN + 300000
TITLE_SLIDE_TOP, TITLE_SLIDE_TITLE_H, TITLE_SLIDE_SUB_H = 1150000, 1100000, \
    700000
LOGO_H, LOGO_GAP, LOGO_BOTTOM = 520000, 520000, 450000


def place(shape, x, y, cx, cy):
    shape.left, shape.top = Emu(x), Emu(y)
    shape.width, shape.height = Emu(cx), Emu(cy)


def is_title(ph):
    # 1 = title, 3 = centred title (title slide)
    return ph.placeholder_format.type in (1, 3)


def set_text_style(ph, size, colour, bold=False):
    """First-level text style of a layout placeholder: left, no bullet"""
    tx_body = ph._element.txBody
    style = etree.fromstring(
        '<a:lstStyle xmlns:a="%s"><a:lvl1pPr marL="0" indent="0" algn="l">'
        '<a:buNone/><a:defRPr sz="%d" b="%d"><a:solidFill>'
        '<a:srgbClr val="%s"/></a:solidFill></a:defRPr></a:lvl1pPr>'
        "</a:lstStyle>" % (A_NS, size, int(bold), colour)
    )
    old = tx_body.find(qn("a:lstStyle"))
    if old is not None:
        tx_body.replace(old, style)
    else:
        tx_body.insert(1, style)


def bar_xml(x, y, cy, width):
    """A McGill red vertical bar, as on the HTML title page"""
    return (
        '<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="900" name="Title bar"/>'
        '<p:cNvCxnSpPr/><p:nvPr userDrawn="1"/></p:nvCxnSpPr><p:spPr>'
        '<a:xfrm><a:off x="%d" y="%d"/><a:ext cx="0" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
        '<a:ln w="%d"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
        "</a:ln></p:spPr></p:cxnSp>"
        % (x, y, cy, width, RED)
    )


def add_logos(layout, paths):
    """Logos at one height, centred along the bottom of a layout"""
    parts = [layout.part.get_or_add_image_part(path) for path in paths]
    widths = [LOGO_H * part.image.size[0] // part.image.size[1]
              for part, _ in parts]
    x = (W - sum(widths) - LOGO_GAP * (len(paths) - 1)) // 2
    y = H - LOGO_BOTTOM - LOGO_H
    for k, (path, (_, rid), cx) in enumerate(zip(paths, parts, widths)):
        pic = CT_Picture.new_pic(910 + k, "Logo %d" % (k + 1), path, rid,
                                 x, y, cx, LOGO_H)
        layout.shapes._spTree.append(pic)
        x += cx + LOGO_GAP


# --- Placeholder geometry and text styles (python-pptx) ---
default = subprocess.run(
    ["quarto", "pandoc", "--print-default-data-file", "reference.pptx"],
    check=True, capture_output=True,
).stdout
prs = Presentation(io.BytesIO(default))

for ph in prs.slide_masters[0].placeholders:
    if is_title(ph):
        place(ph, MARGIN, TITLE_TOP, W - 2 * MARGIN, TITLE_H)
    elif ph.placeholder_format.idx == 1:
        place(ph, MARGIN, BODY_TOP, W - 2 * MARGIN, BODY_H)

for layout in prs.slide_layouts:
    for ph in layout.placeholders:
        idx = ph.placeholder_format.idx
        if layout.name == "Title Slide":
            if is_title(ph):
                place(ph, TEXT_X, TITLE_SLIDE_TOP, W - TEXT_X - MARGIN,
                      TITLE_SLIDE_TITLE_H)
                set_text_style(ph, 3600, INK, bold=True)
            elif ph.placeholder_format.type == 4:  # subtitle
                place(ph, TEXT_X, TITLE_SLIDE_TOP + TITLE_SLIDE_TITLE_H,
                      W - TEXT_X - MARGIN, TITLE_SLIDE_SUB_H)
                set_text_style(ph, 2000, SOFT)
        elif is_title(ph):
            place(ph, MARGIN, TITLE_TOP, W - 2 * MARGIN, TITLE_H)
            if layout.name == "Content with Caption":
                # this layout sets its own smaller title size
                set_text_style(ph, 2600, INK, bold=True)
        elif layout.name == "Title and Content" and idx == 1:
            place(ph, MARGIN, BODY_TOP, W - 2 * MARGIN, BODY_H)
        elif layout.name == "Content with Caption" and idx == 1:
            place(ph, MARGIN, BODY_TOP, W - 2 * MARGIN,
                  NOTE_TOP - BODY_TOP - 100000)
        elif layout.name == "Content with Caption" and idx == 2:
            place(ph, MARGIN, NOTE_TOP, W - 2 * MARGIN, NOTE_H)
            set_text_style(ph, 1400, MUTED)

add_logos(prs.slide_layouts.get_by_name("Title Slide"), LOGOS)

buffer = io.BytesIO()
prs.save(buffer)

# --- Theme, master text styles and rules (direct XML) ---
COLOURS = {
    "dk2": INK, "lt2": "F5F5F5", "accent1": "404040", "accent2": RED,
    "accent3": TEAL, "accent4": "4477AA", "accent5": "228833",
    "accent6": "997700", "hlink": "177DB3", "folHlink": MUTED,
}


def edit_theme(t):
    t = re.sub(r'(<a:(?:major|minor)Font><a:latin typeface=")[^"]*',
               r"\1Arial", t)
    for name, value in COLOURS.items():
        t = re.sub(r"<a:%s>.*?</a:%s>" % (name, name),
                   '<a:%s><a:srgbClr val="%s"/></a:%s>' % (name, value, name),
                   t, flags=re.S)
    return t


def edit_master(t):
    # Titles: left-aligned, bold, 26 pt, anchored at the bottom
    t = re.sub(r'(<p:titleStyle><a:lvl1pPr) algn="ctr"', r'\1 algn="l"', t)
    t = re.sub(r'(<p:titleStyle>.*?<a:defRPr) sz="\d+"',
               r'\1 sz="2600" b="1"', t, flags=re.S)
    t = re.sub(r'(<p:titleStyle>.*?)anchor="ctr"', r'\1anchor="b"', t,
               flags=re.S)
    # Body text: 18 pt, then 16 and 14 pt for deeper levels
    body = re.search(r"<p:bodyStyle>.*?</p:bodyStyle>", t, re.S).group(0)
    new_body = body
    for level, size in ((1, 1800), (2, 1600), (3, 1400), (4, 1400),
                        (5, 1400)):
        new_body = re.sub(r'(<a:lvl%dpPr.*?<a:defRPr) sz="\d+"' % level,
                          r'\1 sz="%d"' % size, new_body, count=1, flags=re.S)
    return t.replace(body, new_body)


def edit_title_slide(t):
    # The red bar beside the title and subtitle, as on the HTML title page
    t = t.replace("<p:sldLayout ", '<p:sldLayout showMasterSp="0" ', 1)
    bar = bar_xml(BAR_X, TITLE_SLIDE_TOP,
                  TITLE_SLIDE_TITLE_H + TITLE_SLIDE_SUB_H, 50800)
    return t.replace("</p:spTree>", bar + "</p:spTree>", 1)


EDITS = {
    "ppt/theme/theme1.xml": edit_theme,
    "ppt/slideMasters/slideMaster1.xml": edit_master,
    "ppt/slideLayouts/slideLayout1.xml": edit_title_slide,
}
with zipfile.ZipFile(buffer) as zin, \
        zipfile.ZipFile(OUTPUT, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in EDITS:
            data = EDITS[item.filename](data.decode("utf8")).encode("utf8")
        zout.writestr(item, data)
print("Saved", OUTPUT)
